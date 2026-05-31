"""
DOCX document loader using python-docx.
"""
import logging
from pathlib import Path
from dataclasses import dataclass
from typing import List

from docx import Document

logger = logging.getLogger(__name__)


class DOCXLoader:
    def load(self, path: str, source: str = "unknown"):
        """Load DOCX and extract text."""
        from ingestion.loaders.pdf_loader import RawDocument

        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f"DOCX not found: {path}")

        doc = Document(str(path))
        full_text = "\n".join(
            para.text.strip() for para in doc.paragraphs if para.text.strip()
        )

        pages = [{"page_num": 1, "text": full_text}]
        logger.info(f"Loaded {path.name}: {len(full_text)} characters")

        return RawDocument(
            filename=path.name,
            source=source,
            page_count=1,
            pages=pages,
        )
